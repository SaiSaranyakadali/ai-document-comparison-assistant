from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt

def create_docx_report(result):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    doc.add_heading("AI Document Comparison Report", 0)
    doc.add_paragraph("Informational comparison generated from the text extracted from two PDF documents.")
    doc.add_paragraph("Important: This report is not legal advice and does not determine legal safety or compliance.")
    for key, label in (("document_a","Document A"),("document_b","Document B")):
        item = result.get(key, {})
        doc.add_heading(label, level=1)
        doc.add_paragraph(f"File: {item.get('file_name','Unknown')}")
        doc.add_paragraph(item.get("summary","Not found in the provided text."))
    for key, label in (("similarities","Similarities"),("differences","Differences")):
        doc.add_heading(label, level=1)
        items = result.get(key, [])
        if items:
            for item in items: doc.add_paragraph(str(item), style="List Bullet")
        else: doc.add_paragraph("No supported points were identified.")
    doc.add_heading("Key Details Comparison", level=1)
    details = result.get("key_details", [])
    if details:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, title in zip(table.rows[0].cells, ["Detail","Document A","Document B","Change / Note"]):
            cell.text = title
        for item in details:
            cells = table.add_row().cells
            for cell, key in zip(cells, ["detail","document_a","document_b","change"]):
                cell.text = str(item.get(key, "Not found in the provided text"))
    else:
        doc.add_paragraph("No key details were returned.")
    doc.add_heading("Overall Conclusion", level=1)
    doc.add_paragraph(result.get("overall_conclusion","Not found in the provided text."))
    doc.add_heading("Limitations", level=1)
    for item in result.get("limitations", []): doc.add_paragraph(str(item), style="List Bullet")
    doc.add_paragraph("The comparison depends on the quality and completeness of the extracted PDF text and the LLM response.")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
